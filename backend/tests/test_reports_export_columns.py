import io
from datetime import date, datetime, timedelta, timezone
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import csv
from sqlalchemy.orm import Session

from app.models.asset import Asset
from app.models.workday_employee import WorkdayEmployee
from app.models.device_eol_reference import DeviceEolReference

EXPECTED_21_COLUMNS = [
    "Asset",
    "Source",
    "DeviceType",
    "Manufacturer",
    "Model",
    "SerialNumber",
    "Operating System",
    "MacAddress",
    "Computer ID",
    "Asset Status",
    "EOL Expiry Date",
    "EOL Status",
    "EOL Days Remaining",
    "Managed By",
    "Assigned UserID",
    "Assigned UserName",
    "Department",
    "Manager UserID",
    "Manager UserName",
    "Job Title",
    "Remarks",
]

SIX_REPORTS = [
    "master_hardware",
    "warranty_compliance_issues",
    "bigfix_assets",
    "jamf_assets",
    "in_stock_inventory",
    "need_validation_assets",
]


def _read_xlsx(content: bytes):
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 1, "XLSX must have at least one header row"
    header = [str(col) if col is not None else "" for col in rows[0]]
    data_rows = rows[1:]
    return header, data_rows


def _read_csv(content: str):
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    assert len(rows) >= 1, "CSV must have at least one header row"
    return rows[0], rows[1:]


def test_exact_21_columns_in_all_six_reports_xlsx(client: TestClient):
    """
    Verify that all six report exports in XLSX format produce the exact 21 columns in the exact order.
    """
    for rep in SIX_REPORTS:
        res = client.get(f"/api/assets/export?report={rep}")
        assert res.status_code == 200, f"Report {rep} failed with status {res.status_code}"
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
        header, rows = _read_xlsx(res.content)
        assert header == EXPECTED_21_COLUMNS, f"Mismatch in report {rep}: {header} != {EXPECTED_21_COLUMNS}"
        assert len(header) == 21


def test_exact_21_columns_in_all_six_reports_csv(client: TestClient):
    """
    Verify that report exports in CSV format (e.g. ?report=...&format=csv) also produce the exact 21 columns in exact order.
    """
    for rep in SIX_REPORTS:
        res = client.get(f"/api/assets/export?report={rep}&format=csv")
        assert res.status_code == 200, f"Report {rep} CSV failed with status {res.status_code}"
        assert "text/csv" in res.headers["content-type"]
        header, rows = _read_csv(res.text)
        assert header == EXPECTED_21_COLUMNS, f"CSV Mismatch in report {rep}: {header} != {EXPECTED_21_COLUMNS}"
        assert len(header) == 21


def test_computer_id_mapping_and_null_preservation(client: TestClient):
    """
    Verify Computer ID mapping:
    - assets.computer_id is mapped to column 9 (index 8)
    - Populated values appear correctly (COMP-1001, COMP-1002, COMP-1004, COMP-1006)
    - NULL values are preserved as blank/empty string without substituting asset_id or asset_number
    """
    res = client.get("/api/assets/export?report=master_hardware")
    assert res.status_code == 200
    header, rows = _read_xlsx(res.content)
    assert header[8] == "Computer ID"

    row_map = {r[0]: r for r in rows}

    # AST-1001 has computer_id="COMP-1001"
    assert row_map["AST-1001"][8] == "COMP-1001"
    assert row_map["AST-1001"][8] != row_map["AST-1001"][0]  # Not asset_number

    # AST-1002 has computer_id="COMP-1002"
    assert row_map["AST-1002"][8] == "COMP-1002"

    # AST-1003 has computer_id=None -> must be blank
    assert (row_map["AST-1003"][8] or "") == ""
    assert row_map["AST-1003"][8] != "AST-1003"
    assert row_map["AST-1003"][8] != "3"

    # AST-1004 has computer_id="COMP-1004"
    assert row_map["AST-1004"][8] == "COMP-1004"

    # AST-1005 has computer_id=None -> must be blank
    assert (row_map["AST-1005"][8] or "") == ""

    # AST-1006 has computer_id="COMP-1006"
    assert row_map["AST-1006"][8] == "COMP-1006"


def test_employee_manager_and_job_title_mapping(client: TestClient):
    """
    Verify:
    - Manager UserID mapped from workday_employees.manager_user_name
    - Manager UserName mapped from workday_employees.manager_name
    - Job Title mapped from workday_employees.job_title
    - Tests:
      1. Matched employee with manager (johndoe -> janemanager, Jane Manager, Senior Systems Engineer)
      2. Matched employee with manager (sarahconnor -> robertboss, Robert Boss, Lead Security Analyst)
      3. Matched employee with NULL manager (alexturner -> blank, blank, Growth Specialist)
      4. Assigned user not found in Workday (admin_local -> blank, blank, blank)
      5. Unassigned asset (None -> blank, blank, blank)
    """
    res = client.get("/api/assets/export?report=master_hardware")
    assert res.status_code == 200
    header, rows = _read_xlsx(res.content)
    row_map = {r[0]: r for r in rows}

    # 1. John Doe: Full manager and job title
    ast1 = row_map["AST-1001"]
    assert ast1[16] == "Engineering"  # Department
    assert ast1[17] == "janemanager"  # Manager UserID
    assert ast1[18] == "Jane Manager"  # Manager UserName
    assert ast1[19] == "Senior Systems Engineer"  # Job Title
    assert ast1[20] == "Primary developer laptop"  # Remarks

    # 2. Sarah Connor: Full manager and job title
    ast2 = row_map["AST-1002"]
    assert ast2[16] == "Security"
    assert ast2[17] == "robertboss"
    assert ast2[18] == "Robert Boss"
    assert ast2[19] == "Lead Security Analyst"
    assert ast2[20] == "Needs security check"

    # 3. Unassigned asset: All employee columns blank
    ast3 = row_map["AST-1003"]
    assert (ast3[14] or "") == ""  # Assigned UserID
    assert (ast3[15] or "") == ""  # Assigned UserName
    assert (ast3[16] or "") == ""  # Department
    assert (ast3[17] or "") == ""  # Manager UserID
    assert (ast3[18] or "") == ""  # Manager UserName
    assert (ast3[19] or "") == ""  # Job Title

    # 4. admin_local (Not in Workday): All employee columns blank, no crash or fake values
    ast4 = row_map["AST-1004"]
    assert ast4[14] == "admin_local"
    assert (ast4[15] or "") == ""
    assert (ast4[16] or "") == ""
    assert (ast4[17] or "") == ""
    assert (ast4[18] or "") == ""
    assert (ast4[19] or "") == ""

    # 5. alexturner (In Workday, manager fields are None): Job title populated, manager fields blank
    ast5 = row_map["AST-1005"]
    assert ast5[15] == "Alex Turner"
    assert ast5[16] == "Marketing"
    assert (ast5[17] or "") == ""  # manager_user_name is None
    assert (ast5[18] or "") == ""  # manager_name is None
    assert ast5[19] == "Growth Specialist"


def test_asset_deduplication_in_reports(client: TestClient, db_session: Session):
    """
    Ensure the outer join to Workday does not duplicate assets even if employee records exist.
    """
    res = client.get("/api/assets/export?report=master_hardware")
    assert res.status_code == 200
    _, rows = _read_xlsx(res.content)
    asset_numbers = [r[0] for r in rows]
    assert len(asset_numbers) == len(set(asset_numbers)), "Assets must be unique in export"
    assert len(rows) == 6
